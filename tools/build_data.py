#!/usr/bin/env python3
"""Build the CCR app data.json (original compact schema) from Paul's Excel files.

Usage:
  python3 tools/build_data.py --input /path/to/excel_folder [--base data_base.json] [--current data.json] [--out data.json]

* Files are found by filename pattern (see PATTERNS); when several match, the one with the newest
  date in its name (dd.mm.yyyy / yyyy.mm.dd) or highest R-number wins, then the newest mtime.
* Sections without an Excel source (lr, drums, mat, mrf, sf, rv, sg, panels) are carried forward from
  the current data.json (if it has the right schema) or data_base.json.
* Cables already in the app keep their From/To text (f, fr, o, tr) unchanged.
* Standard references (GOST/SP/PUE/...) are removed from all text.
* Output is validated against data_base.json's key structure; on failure nothing is written, exit 1.
"""
import argparse, datetime, glob, json, os, re, sys, time

SITES = {'ccr': (('G0-D1810BX0', 'G0-D1801FX0'), 'D1810|D1801'),
         'd48': (('G0-D4810BX0', 'G0-D4801FX0'), 'D4810|D4801')}
SITE = 'ccr'
CCR, SITE_RX = SITES['ccr']   # building tags of the selected site (name kept for history)
PATTERNS = {
    'iwp':    ['*IWP_Status*.xls*'],
    'cables': ['*Cable_Tracker*.xls*'],
    'rhi':    ['*Internal_Punch*.xls*'],
    'master': ['*Master_Punch_List*.xls*'],
    'equip':  ['*Equipment_Control*.xls*'],
    'rawp':   ['*RAWP*.xls*'],
    'docs':   ['*documents register*.xls*', '*Documents_Register*.xls*'],
    'ptq':    ['*Technical Query Register*.xls*'],
    'ctq':    ['*CCR*TQ_Register*.xls*', '*CTQ*Register*.xls*'],
}
SHEETS = {'iwp': 'EL_IWP_Status', 'cables': 'Единый кабельный журнал', 'rhi': 'Punch_List',
          'master': 'Log of Mechanical Progress', 'equip': 'E&I_Equipment', 'rawp': 'Cable_list',
          'docs': None, 'ptq': 'TQ Registry', 'ctq': 'TQ Register'}
LABEL = {'iwp': 'IWP status', 'cables': 'Md cable tracker', 'rhi': 'RHI internal punch', 'master': 'Master punch',
         'equip': 'E&I equipment control', 'rawp': 'Cable RAWP', 'docs': 'Documents register',
         'ptq': 'TQ register', 'ctq': 'CTQ register'}


def log(*a):
    print(*a, flush=True)


# ---------------------------------------------------------------- helpers
STD_RX = re.compile(r'(?<![\w\-])(ГОСТ|GOST|СП|SP|ПУЭ|PUE|СНиП|SNiP|IEC|МЭК|ISO|NFPA|ВСН|НПБ|EN)\s?(Р|R)?\s?\d+(?:[.\-–/:]\d+)*', re.I)
STD_N = [0]


def scrub(s):
    def rep(m):
        STD_N[0] += 1
        return 'норм. док.' if re.match(r'[А-Яа-я]', m.group(1)) else 'standard'
    return STD_RX.sub(rep, s)


def scrub_all(o):
    if isinstance(o, str):
        return scrub(o) if STD_RX.search(o) else o
    if isinstance(o, list):
        return [scrub_all(x) for x in o]
    if isinstance(o, dict):
        return {k: scrub_all(v) for k, v in o.items()}
    return o


def S(v):
    """clean string: None/'-'/0-placeholders -> '', collapse whitespace"""
    if v is None:
        return ''
    if isinstance(v, bool):
        return ''
    if isinstance(v, float):
        v = int(v) if v == int(v) else v
    s = ' '.join(str(v).split())
    return '' if s in ('-', 'None', '#N/A', '#VALUE!') else s


def num(v):
    if v is None or v == '' or isinstance(v, bool):
        return 0
    try:
        f = float(str(v).replace(',', '.'))
    except ValueError:
        return 0
    return int(f) if f == int(f) else round(f, 3)


def acc_s(v):
    try:
        return str(int(float(v)))
    except (TypeError, ValueError):
        return S(v)


def xdate(v, fmt='%d.%m.%Y'):
    if v is None or v == '':
        return ''
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime(fmt)
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        if 20000 < v < 80000:
            return (datetime.date(1899, 12, 30) + datetime.timedelta(days=int(v))).strftime(fmt)
        return ''
    s = S(v)
    m = re.match(r'(\d{4})-(\d{2})-(\d{2})', s)
    if m:
        return datetime.date(int(m[1]), int(m[2]), int(m[3])).strftime(fmt)
    return s


def col(r, i):
    return r[i] if i is not None and i < len(r) else None


class Tab:
    """index lookup that appends new values to a tabs list"""
    def __init__(self, lst):
        self.l = lst
        self.ix = {}
        for i, v in enumerate(lst):
            self.ix.setdefault(v, i)

    def __call__(self, v):
        if v not in self.ix:
            self.ix[v] = len(self.l)
            self.l.append(v)
        return self.ix[v]


# ---------------------------------------------------------------- file discovery / reading
def file_key(path):
    n = os.path.basename(path)
    m = re.search(r'(\d{2})\.(\d{2})\.(\d{4})', n)
    d = None
    if m:
        try:
            d = datetime.date(int(m[3]), int(m[2]), int(m[1]))
        except ValueError:
            d = None
    m2 = re.search(r'(\d{4})\.(\d{2})\.(\d{2})', n)
    if not d and m2:
        try:
            d = datetime.date(int(m2[1]), int(m2[2]), int(m2[3]))
        except ValueError:
            d = None
    rn = re.search(r'_R(\d+)', n)
    return (d or datetime.date(1900, 1, 1), int(rn[1]) if rn else 0, os.path.getmtime(path)), d


def _norm_name(s):
    return re.sub(r'[\s_]+', ' ', s).lower()


def find_files(folder):
    """pattern match ignores upper/lower case and treats spaces and underscores the same"""
    import fnmatch
    names = [f for f in glob.glob(os.path.join(folder, '**', '*'), recursive=True)
             if os.path.isfile(f) and not os.path.basename(f).startswith('~$')]
    out = {}
    for k, pats in PATTERNS.items():
        c = set()
        for p in pats:
            c.update(f for f in names if fnmatch.fnmatch(_norm_name(os.path.basename(f)), _norm_name(p)))
        if c:
            out[k] = max(c, key=lambda f: file_key(f)[0])
    return out


def read_sheet(path, sheet):
    t = time.time()
    if path.lower().endswith('.xlsb'):
        from pyxlsb import open_workbook
        with open_workbook(path) as wb:
            name = sheet or wb.sheets[0]
            with wb.get_sheet(name) as ws:
                rows = [[c.v for c in r] for r in ws.rows()]
    else:
        import openpyxl, warnings
        warnings.filterwarnings('ignore')
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb[sheet] if sheet else wb.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
    while rows and not any(v not in (None, '') for v in rows[-1]):
        rows.pop()
    log(f'  read {os.path.basename(path)} [{sheet or "first sheet"}]: {len(rows)} rows, {time.time() - t:.0f}s')
    return rows


# ---------------------------------------------------------------- section builders
def build_items(rows, prev, T):
    """IWP status -> items (IK positional rows). Column numbers = 0-based sheet columns."""
    IK = prev['IK']
    tb = {k: Tab(T[k]) for k in ('grp', 'map', 'unit', 'desc', 'dwg', 'sub', 'sys', 'co', 'resp', 'it')}
    lower = {k: {v.lower(): v for v in T[k]} for k in ('map', 'grp')}
    prev_order = {r[0]: i for i, r in enumerate(prev['items'])}
    new = []
    for r in rows[5:]:
        if len(r) < 47 or S(col(r, 2)) not in CCR:
            continue
        acc = acc_s(col(r, 34))
        if not acc:
            continue
        mp = S(col(r, 1))
        mp = lower['map'].get(mp.lower(), mp)
        grp = 'Cable pulling' if re.search(r'cable laying', mp, re.I) else 'Cable termination' if re.search(r'cable termination', mp, re.I) else mp
        grp = lower['grp'].get(grp.lower(), grp)
        dwg = S(col(r, 5))
        m = re.search(r'\)-([A-Z]{2}-[A-Z0-9]{4})', dwg)
        sysv = m.group(1) if m else S(col(r, 4))
        q = col(r, 42)
        q = num(q) if q not in (None, '') else num(col(r, 21))
        co = S(col(r, 46))
        co = co if co and co != '0' else 'Not assigned'
        o = {'acc': acc, 'grp': tb['grp'](grp), 'map': tb['map'](mp), 'unit': tb['unit'](S(col(r, 22)) or 'pc'),
             'tag': S(col(r, 14)), 'uid': S(col(r, 11)), 'desc': tb['desc'](S(col(r, 20))[:140]), 'dwg': tb['dwg'](dwg),
             'sub': tb['sub'](S(col(r, 3))), 'sys': tb['sys'](sysv), 'q': q, 'f': num(col(r, 43)),
             'co': tb['co'](co), 'resp': tb['resp'](S(col(r, 37))), 'notes': S(col(r, 33)), 'tq': S(col(r, 32)),
             'it': tb['it'](S(col(r, 38))[:90]), 'ctag': S(col(r, 13))[:60]}
        new.append([o[k] for k in IK])
    seen, items = set(), []
    for row in sorted(new, key=lambda x: prev_order.get(x[0], 10 ** 9)):
        if row[0] in seen:
            continue
        seen.add(row[0])
        items.append(row)
    return items


def enrich_items_equipment(items, rows, IK):
    """E&I equipment control -> RFI status written into item notes (matched by unique tag)."""
    iu, inn = IK.index('uid'), IK.index('notes')
    rfi = {}
    for r in rows[2:]:
        if len(r) < 26 or S(col(r, 1)) not in CCR:
            continue
        u, no = S(col(r, 18)), S(col(r, 23))
        if u and no:
            rfi[u] = ' · '.join(x for x in ('RFI ' + no, S(col(r, 24)), xdate(col(r, 25))) if x)
    n = 0
    for it in items:
        v = rfi.get(it[iu])
        if v and v not in it[inn]:
            it[inn] = (it[inn] + ' · ' if it[inn] else '') + v
            n += 1
    return n, len(rfi)


def tracker_index(rows):
    by = {}
    for r in rows[5:]:
        t = S(col(r, 21))
        if t:
            by.setdefault(t, []).append(r)
    return by


def pick_row(cands, prev=None, CK=None):
    """several tracker rows can share one cable number: prefer the row that matches the existing cable"""
    if len(cands) == 1:
        return cands[0]
    pty = pd_ = ''
    if prev is not None:
        pty, pd_ = prev[CK.index('ty')], prev[CK.index('d')]
    pt = pd_.split(' ')[0] if pd_ else ''
    for test in (lambda r: pty and S(col(r, 26)) == pty and not S(col(r, 35)),
                 lambda r: pty and S(col(r, 26)) == pty,
                 lambda r: pt and S(col(r, 14)) == pt and not S(col(r, 35)),
                 lambda r: S(col(r, 14)) in CCR and not S(col(r, 35)),
                 lambda r: S(col(r, 75)) == 'RHI' and not S(col(r, 35)),
                 lambda r: not S(col(r, 35))):
        c = [r for r in cands if test(r)]
        if c:
            return c[0]
    return cands[0]


def room(s):
    m = re.search(r'Room\s+([0-9A-Za-zА-Яа-я.\-]+)', s)
    return m.group(1) if m else ''


def hold_text(r):
    """HOLD column; if empty, the 'counted in another project' column (deleted / duplicate ...)"""
    b, a = S(col(r, 36)), S(col(r, 35))
    if b:
        return b
    if a:
        return 'в другом проекте' if a.lower() in ('да', 'yes') else a
    return ''


def drum(r):
    return S(col(r, 48)) or S(col(r, 64)) or S(col(r, 60))


def build_cables(rows, prev, items, rawp_rows):
    CK = prev['CK']
    K = {k: i for i, k in enumerate(CK)}
    by = tracker_index(rows)
    cables = [list(c) for c in prev['cables']]
    have = {c[0] for c in cables}
    upd = 0
    for c in cables:
        if c[K['t']] not in by:
            continue
        r = pick_row(by[c[K['t']]], c, CK)
        upd += 1
        doc = S(col(r, 0))
        c[K['d']] = (re.sub(r'^GLE-\((.*?)\)-', r'\1 ', doc) + (' rev ' + S(col(r, 5)) if S(col(r, 5)) else '')) or c[K['d']]
        # f, fr, o, tr are kept from the existing data (user rule)
        for k, i in (('ty', 26), ('c', 27), ('tp', 50), ('tt', 55)):
            c[K[k]] = S(col(r, i))
        c[K['dr']] = drum(r)
        ln = num(col(r, 32))
        c[K['ln']] = str(ln) if ln else c[K['ln']]
        c[K['pd']] = xdate(col(r, 49), '%d.%m.%y')
        c[K['tf']] = xdate(col(r, 52), '%d.%m.%y')
        c[K['tto']] = xdate(col(r, 53), '%d.%m.%y')
        c[K['h']] = hold_text(r)
        if S(col(r, 75)) == 'RHI':
            c[K['co']] = 'RHI'
    added = 0
    for t, cands in by.items():
        if t in have:
            continue
        r = pick_row(cands)
        if not (S(col(r, 14)) in CCR or S(col(r, 24)) in CCR or S(col(r, 25)) in CCR):
            continue
        if S(col(r, 35)):          # deleted / duplicated / in another project -> not added
            continue
        f, o = S(col(r, 22)), S(col(r, 23))
        doc = S(col(r, 0))
        row = [''] * len(CK)
        vals = {'t': t, 'd': re.sub(r'^GLE-\((.*?)\)-', r'\1 ', doc) + (' rev ' + S(col(r, 5)) if S(col(r, 5)) else ''),
                'f': f, 'fr': room(f), 'o': o, 'tr': room(o), 'ty': S(col(r, 26)), 'c': S(col(r, 27)),
                'ln': str(num(col(r, 32))) if num(col(r, 32)) else '', 'dr': drum(r), 'tp': S(col(r, 50)),
                'pd': xdate(col(r, 49), '%d.%m.%y'), 'tt': S(col(r, 55)), 'tf': xdate(col(r, 52), '%d.%m.%y'),
                'tto': xdate(col(r, 53), '%d.%m.%y'), 'co': 'RHI' if S(col(r, 75)) == 'RHI' else '',
                'tqs': S(col(r, 18)), 'h': hold_text(r), 'src': 'Cable tracker'}
        for k, v in vals.items():
            row[K[k]] = v
        for k in ('lq', 'lf', 'tq', 'tfa'):
            row[K[k]] = ''
        cables.append(row)
        have.add(t)
        added += 1
    # IWP access ids / qty / fact per cable
    IK = prev['IK']
    ia = {k: IK.index(k) for k in ('acc', 'grp', 'tag', 'q', 'f')}
    grp = prev['tabs']['grp']
    pull, term = {}, {}
    for it in items:
        g = grp[it[ia['grp']]]
        d = pull if re.search('pull', g, re.I) else term if re.search('termination', g, re.I) else None
        if d is not None and it[ia['tag']]:
            d.setdefault(it[ia['tag']], []).append(it)
    ncab = {}
    for c in cables:
        ncab[c[K['t']]] = ncab.get(c[K['t']], 0) + 1

    def fill(c, lst, ka, kq, kf):
        if not lst:
            return
        same = [x for x in lst if x[ia['acc']] == c[K[ka]]]
        if same:                      # keep the Access ID already linked, refresh its qty / fact
            lst = same[:1]
        elif ncab[c[K['t']]] > 1 or not re.search(r'\d', c[K['t']]) or 'NO TAG' in c[K['t']].upper():
            return                    # generic / repeated cable number: do not guess
        c[K[ka]] = lst[0][ia['acc']]
        c[K[kq]] = num(sum(x[ia['q']] for x in lst))
        c[K[kf]] = num(sum(x[ia['f']] for x in lst))
    for c in cables:
        fill(c, pull.get(c[K['t']]), 'la', 'lq', 'lf')
        fill(c, term.get(c[K['t']]), 'ta', 'tq', 'tfa')
    # RAWP: fill missing laying / termination dates and add RFI info to the comment
    rw = 0
    if rawp_rows:
        ix = {c[K['t']]: c for c in cables}
        for r in rawp_rows[3:]:
            c = ix.get(S(col(r, 8)))
            if not c:
                continue
            ld, td = xdate(col(r, 17), '%d.%m.%y'), xdate(col(r, 24), '%d.%m.%y')
            if num(col(r, 13)) >= 1 and ld and not c[K['pd']]:
                c[K['pd']] = ld
            if num(col(r, 20)) >= 1 and td and not c[K['tf']] and not c[K['tto']]:
                c[K['tf']] = td
            info = []
            if S(col(r, 15)):
                info.append(f'Laying RFI {S(col(r, 15))} {S(col(r, 16))}'.strip())
            if S(col(r, 22)):
                info.append(f'Term. RFI {S(col(r, 22))} {S(col(r, 23))}'.strip())
            cm = re.sub(r'\s*\[RAWP:[^\]]*\]', '', c[K['cm']] or '')
            if info:
                cm = (cm + ' ' if cm else '') + '[RAWP: ' + '; '.join(info) + ']'
                rw += 1
            c[K['cm']] = cm
    return cables, upd, added, rw


def update_panels(panels, cables, prev_n, CK):
    """new cables (appended after prev_n) are linked to panels whose tag appears in To (IN) / From (OUT)."""
    it, io, ifr = CK.index('t'), CK.index('o'), CK.index('f')
    pans = [p for p in panels if len(p[0]) >= 8]
    n = 0
    for c in cables[prev_n:]:
        for p in pans:
            if p[0] in (c[io] or ''):
                p[3].append(['IN', c[it]]); n += 1
            if p[0] in (c[ifr] or ''):
                p[3].append(['OUT', c[it]]); n += 1
    return n


CYR = str.maketrans('АВЕКМНОРСТХ', 'ABEKMHOPCTX')
TAG_RX = re.compile(r'G[01]-[A-Z0-9]{4,12}(?:-[A-Z0-9()]{1,12}){1,6}')


def build_panels(cables, CK):
    """site without a panel list yet (D48): every equipment tag named in From / To becomes a panel."""
    it, io, ifr = CK.index('t'), CK.index('o'), CK.index('f')
    pans = {}
    for c in cables:
        for side, txt in (('OUT', c[ifr] or ''), ('IN', c[io] or '')):
            tags = [t.rstrip('-') for t in TAG_RX.findall(txt.translate(CYR))]
            tags = [t for t in tags if t != c[it] and not t.endswith('-WP01')]
            if not tags:
                continue
            tag = max(tags, key=len)
            p = pans.setdefault(tag, [tag, txt[:160], room(txt), []])
            p[3].append([side, c[it]])
    out = [p for p in pans.values() if any(x in p[0] for x in CCR) or len(p[3]) >= 2 or re.search(r'-D48-|D48', p[0])]
    out.sort(key=lambda p: p[0])
    return out


def build_sg(cables, CK, panels):
    """D48: system block diagrams from the cable list. Node = panel / box (tag found in From / To);
    devices with one cable are grouped per panel. Columns follow the cable direction (from -> to)."""
    K = {k: i for i, k in enumerate(CK)}
    pdesc = {p[0]: p[1] for p in panels}
    dead = re.compile('удален|дубликат|другом проекте', re.I)

    def end(txt):
        tags = [t.rstrip('-') for t in TAG_RX.findall((txt or '').translate(CYR))]
        tags = [t for t in tags if not t.endswith('-WP01')]
        return max(tags, key=len) if tags else (txt or '-').strip()[:50] or '-'
    bysys = {}
    for i, c in enumerate(cables):
        m = re.search(r'\b([A-Z]{2}-[A-Z0-9]{4})-', c[K['d']] or '')
        if not m or dead.search(c[K['h']] or ''):
            continue
        a, b = end(c[K['f']]), end(c[K['o']])
        if a == b:
            continue
        bysys.setdefault(m.group(1), []).append((i, a, b, c[K['f']], c[K['o']]))
    sg = {}
    for sy, lst in sorted(bysys.items()):
        deg = {}
        txt = {}
        for i, a, b, fa, fb in lst:
            deg[a] = deg.get(a, 0) + 1; deg[b] = deg.get(b, 0) + 1
            txt.setdefault(a, fa); txt.setdefault(b, fb)
        hub = {t for t, n in deg.items() if n >= 2 or t in pdesc}
        node, grp = {}, {}

        def nid(t, other, side):
            if t in hub or other not in hub:
                return node.setdefault(t, ('T', t))
            k = '#' + other + '|' + side
            grp.setdefault(k, [])
            if t not in grp[k]:
                grp[k].append(t)
            return node.setdefault(k, ('G', k))
        edges = {}
        for i, a, b, fa, fb in lst:
            x, y = nid(a, b, 'out'), nid(b, a, 'in')
            edges.setdefault((x[1], y[1]), []).append(i)
        keys = list(dict.fromkeys([k for e in edges for k in e]))
        # columns: longest path from sources (cycles cut by visit order)
        succ = {}
        for (x, y) in edges:
            succ.setdefault(x, set()).add(y)
        col = {k: 0 for k in keys}
        for _ in range(6):
            ch = False
            for (x, y) in edges:
                if col[y] < col[x] + 1 and col[x] < 5:
                    col[y] = col[x] + 1; ch = True
            if not ch:
                break
        rows = {}
        order = sorted(keys, key=lambda k: (col[k], k.startswith('#'), k))
        n_out, ix = [], {}
        for k in order:
            r = rows.get(col[k], 0); rows[col[k]] = r + 1
            ix[k] = len(n_out)
            if k.startswith('#'):
                mem = grp[k]
                n_out.append([k, f'{len(mem)} devices / boxes', '; '.join(m[:40] for m in mem[:3]), col[k], r, 1, mem])
            else:
                n_out.append([k, k, (pdesc.get(k) or txt.get(k) or '')[:120], col[k], r, 0, []])
        e_out = [[ix[x], ix[y], sorted(set(cs))] for (x, y), cs in edges.items()]
        if len(n_out) >= 2:
            sg[sy] = {'n': n_out, 'e': e_out}
    return sg


def st_map(v):
    s = S(v).lower()
    return 'Closed' if re.search(r'закр|clos|отмен|cancel', s) else 'Open'


def build_punch(mrows, rrows, prev):
    P = json.loads(json.dumps(prev['punch']))
    K, PT = P['K'], P['tabs']
    ru_en = {}
    for ru, en in (('subn', 'subnE'), ('disc', 'discE'), ('stg', 'stgE'), ('sec', 'secE')):
        ru_en[ru] = dict(zip(PT[ru], PT[en]))
    tb = {k: Tab(PT[k]) for k in PT}

    def tx(k, v):
        i = tb[k](v)
        if k in ru_en:
            en = k + 'E'
            while len(PT[en]) < len(PT[k]):
                PT[en].append(ru_en[k].get(PT[k][len(PT[en])], PT[k][len(PT[en])]))
        return i
    prev_tage = {r[K.index('no')]: r[K.index('tage')] for r in prev['punch']['rows']}
    rows = []
    ino = K.index('no')
    if not mrows:      # no new master punch file: keep the client rows already in the app
        rows = [list(r) for r in prev['punch']['rows'] if not str(r[ino]).startswith('RHI-')]
    for r in (mrows or [])[1:]:
        tit = S(col(r, 1))
        if not tit or not re.search(SITE_RX, tit) or not S(col(r, 0)) or S(col(r, 0)) == '#':
            continue
        no = acc_s(col(r, 0))
        o = {'no': no, 'tit': tx('tit', (tit + ' ' + S(col(r, 2))).strip() if S(col(r, 2)) else tit + ' '),
             'sys': tx('sys', (S(col(r, 3)) + ' ' + S(col(r, 4))).strip()), 'sub': tx('sub', S(col(r, 5))),
             'subn': tx('subn', S(col(r, 6))), 'pri': S(col(r, 7)), 'disc': tx('disc', S(col(r, 8))),
             'tag': S(col(r, 9)), 'tagd': S(col(r, 10)), 'sec': tx('sec', S(col(r, 11))), 'dru': S(col(r, 12)),
             'den': S(col(r, 13)), 'stg': tx('stg', S(col(r, 14))), 'idate': xdate(col(r, 15)),
             'delay': num(col(r, 16)) if S(col(r, 16)) else '', 'cat': S(col(r, 18)), 'by': tx('by', S(col(r, 19))),
             'org': tx('org', S(col(r, 20))), 'epc': tx('epc', S(col(r, 21))), 'resp': tx('resp', S(col(r, 22))),
             'st': st_map(col(r, 23)), 'cdate': xdate(col(r, 24)), 'notes': S(col(r, 25)), 'reason': S(col(r, 26)),
             'dept': tx('dept', S(col(r, 28))), 'insp': S(col(r, 30)), 'act': S(col(r, 31)), 'rem': S(col(r, 32)),
             'tage': prev_tage.get(no, '')}
        rows.append([o[k] for k in K])
    nm = len(rows)
    # internal RHI punch register -> rows numbered RHI-<n>, responsible RHI
    disc = {'EL': 'Электротехническое оборудование', 'IN': 'Оборудование КИПиА', 'TE': 'Системы связи',
            'FA': 'Оборудование АСУ', 'HV': 'Системы отопления, вентиляции и кондиционирования'}
    sub_info = {}
    for r in rows:
        sub_info.setdefault(PT['sub'][r[K.index('sub')]], (r[K.index('sys')], r[K.index('subn')]))
    tit_for = {t: next((i for i, v in enumerate(PT['tit']) if v.startswith(t)), None) for t in CCR}
    today = datetime.date.today()
    if not rrows:      # no new RHI punch file: keep the RHI rows already in the app
        rows += [list(r) for r in prev['punch']['rows'] if str(r[ino]).startswith('RHI-')]
    for r in (rrows or [])[1:]:
        t = S(col(r, 1))
        if t not in CCR or not S(col(r, 0)):
            continue
        sub = S(col(r, 2))
        sy, sn = sub_info.get(sub, (tx('sys', ''), tx('subn', '')))
        idate = xdate(col(r, 12))
        try:
            delay = (today - datetime.datetime.strptime(idate, '%d.%m.%Y').date()).days
        except ValueError:
            delay = ''
        stv = S(col(r, 13))
        notes = ' · '.join(x for x in ('RFI ' + S(col(r, 4)) if S(col(r, 4)) else '', S(col(r, 5)),
                                       'Отменен' if re.search('отмен', stv, re.I) else '', S(col(r, 18))) if x)
        o = {'no': 'RHI-' + acc_s(col(r, 0)), 'tit': tit_for[t] if tit_for[t] is not None else tx('tit', t),
             'sys': sy, 'sub': tx('sub', sub), 'subn': sn, 'pri': S(col(r, 20)),
             'disc': tx('disc', disc.get(S(col(r, 3)), 'АСУ и КИП')), 'tag': S(col(r, 6)), 'tagd': S(col(r, 7)),
             'sec': tx('sec', S(col(r, 8))), 'dru': S(col(r, 9)), 'den': '', 'stg': tx('stg', 'СМР'),
             'idate': idate, 'delay': delay, 'cat': S(col(r, 10)), 'by': tx('by', S(col(r, 11))),
             'org': tx('org', 'РХИ (внутренний)'), 'epc': tx('epc', 'РХИ'), 'resp': tx('resp', 'РХИ'),
             'st': st_map(stv), 'cdate': xdate(col(r, 14)), 'notes': notes, 'reason': '', 'dept': tx('dept', 'E&I'),
             'insp': S(col(r, 21)), 'act': S(col(r, 17)), 'rem': 'RHI internal punch', 'tage': ''}
        rows.append([o[k] for k in K])
    P['rows'] = rows
    return P, nm, len(rows) - nm


def build_reg(rows, prev):
    old = {r[0]: r for r in prev['reg']}
    best = {}
    for r in rows[2:]:
        if len(r) < 30 or S(col(r, 8)) not in CCR or not S(col(r, 16)):
            continue
        no = S(col(r, 16))
        d = col(r, 29)
        key = d if isinstance(d, datetime.datetime) else datetime.datetime(1900, 1, 1)
        if no not in best or key >= best[no][0]:
            best[no] = (key, r)
    out = []
    for no, (_, r) in best.items():
        o = old.get(no, [''] * 9)
        sysv = '-'.join(x for x in (S(col(r, 12)), S(col(r, 13))) if x)
        out.append([no, sysv or o[1], S(col(r, 14)), S(col(r, 17)) or o[3], S(col(r, 18)) or o[4],
                    S(col(r, 28)) or o[5], S(col(r, 27)) or o[6], xdate(col(r, 29)) or o[7], S(col(r, 11)) or o[8]])
    order = {r[0]: i for i, r in enumerate(prev['reg'])}
    out.sort(key=lambda x: (order.get(x[0], 10 ** 9), x[0]))
    return out


def build_ptq(rows, prev):
    best = {}
    for r in rows[1:]:
        no = S(col(r, 4))
        if not no or not any(x in S(col(r, c)) for c in (4, 16, 17) for x in SITE_RX.split('|')):
            continue
        best[no] = r          # later revision rows overwrite earlier ones
    out = []
    for no, r in best.items():
        ep, ow = S(col(r, 27)), S(col(r, 34))
        st = 'Open' if not ep else 'Waiting owner' if ow.lower().startswith('declined') else 'Closed'
        out.append([no, S(col(r, 5)), xdate(col(r, 8)), S(col(r, 14)), S(col(r, 16)), S(col(r, 17)), S(col(r, 19)),
                    S(col(r, 20)), S(col(r, 21))[:500], ep, xdate(col(r, 24)), S(col(r, 29))[:300], ow,
                    S(col(r, 31)), S(col(r, 39)), S(col(r, 42)), S(col(r, 40))[:300], st])
    order = {r[0]: i for i, r in enumerate(prev['tq']['ptq'])}
    out.sort(key=lambda x: (order.get(x[0], 10 ** 9), x[0]))
    return out


def build_ctq(rows, prev):
    hdr = next((i for i, r in enumerate(rows[:15]) if S(col(r, 0)).lower().startswith('tq no')), 3)
    reg = {}
    for r in rows[hdr + 1:]:
        no = S(col(r, 0))
        if re.match(r'CTQ-\d+', no):
            reg[no] = [no] + [xdate(col(r, i)) if i == 17 else S(col(r, i)) for i in (1, 2, 3, 4, 5, 6, 7, 9, 10, 11, 12, 13, 14, 16, 17, 18)]
    out, seen = [], set()
    for r in prev['tq']['ctq']:
        out.append(reg.get(r[0], r))
        seen.add(r[0])
    out += [v for k, v in reg.items() if k not in seen]
    return out, len(reg)


# ---------------------------------------------------------------- validation
def validate(out, base):
    err = []
    if set(out) != set(base):
        err.append(f'top-level keys differ: missing {set(base) - set(out)}, extra {set(out) - set(base)}')
        return err
    for k in base:
        if type(out[k]) is not type(base[k]):
            err.append(f'{k}: type {type(out[k]).__name__} != {type(base[k]).__name__}')
    if err:
        return err
    if set(out['meta']) != set(base['meta']):
        err.append('meta keys differ')
    if set(out['tabs']) != set(base['tabs']):
        err.append('tabs keys differ')
    if out['IK'] != base['IK'] or out['CK'] != base['CK']:
        err.append('IK/CK differ from base')
    T = out['tabs']
    tabk = {'grp', 'map', 'unit', 'desc', 'dwg', 'sub', 'sys', 'co', 'resp', 'it'}
    for n, r in enumerate(out['items']):
        if not isinstance(r, list) or len(r) != len(out['IK']):
            err.append(f'items[{n}] bad length'); break
        for j, k in enumerate(out['IK']):
            if k in tabk and not (isinstance(r[j], int) and 0 <= r[j] < len(T[k])):
                err.append(f'items[{n}].{k} index out of range'); break
            if k in ('q', 'f') and not isinstance(r[j], (int, float)):
                err.append(f'items[{n}].{k} not numeric'); break
    for n, r in enumerate(out['cables']):
        if not isinstance(r, list) or len(r) != len(out['CK']) or not r[0]:
            err.append(f'cables[{n}] bad row'); break
    for n, p in enumerate(out['panels']):
        if not (isinstance(p, list) and len(p) == 4 and isinstance(p[3], list)):
            err.append(f'panels[{n}] bad row'); break
    P, B = out['punch'], base['punch']
    if set(P) != set(B) or P['K'] != B['K'] or set(P['tabs']) != set(B['tabs']):
        err.append('punch structure differs')
    else:
        idx = [j for j, k in enumerate(P['K']) if k in P['tabs']]
        for n, r in enumerate(P['rows']):
            if len(r) != len(P['K']) or any(not (isinstance(r[j], int) and 0 <= r[j] < len(P['tabs'][P['K'][j]])) for j in idx):
                err.append(f'punch.rows[{n}] bad row'); break
            if r[P['K'].index('st')] not in ('Open', 'Closed'):
                err.append(f'punch.rows[{n}] bad status'); break
        for ru in ('subn', 'disc', 'stg', 'sec'):
            if len(P['tabs'][ru]) != len(P['tabs'][ru + 'E']):
                err.append(f'punch.tabs.{ru}E length differs from {ru}')
    if set(out['tq']) != set(base['tq']):
        err.append('tq keys differ')
    else:
        if any(len(r) != 17 for r in out['tq']['ctq']):
            err.append('tq.ctq row length != 17')
        if any(len(r) != 18 for r in out['tq']['ptq']):
            err.append('tq.ptq row length != 18')
    if any(len(r) != 9 for r in out['reg']):
        err.append('reg row length != 9')
    nc = len(out['cables'])
    for s, g in out['sg'].items():
        if not isinstance(g, dict) or set(g) != {'n', 'e'}:
            err.append(f'sg.{s} bad'); break
        if any(not (isinstance(i, int) and 0 <= i < nc) for e in g['e'] for i in e[2]):
            err.append(f'sg.{s} cable index out of range'); break
    for k in ('items', 'cables', 'panels'):
        if not out[k]:
            err.append(f'{k} is empty')
    if not out['punch']['rows']:
        err.append('punch.rows is empty')
    if not out['tq']['ctq'] and SITE == 'ccr':
        err.append('tq.ctq is empty')
    return err


def ok_schema(d, base):
    try:
        return not validate(d, base)
    except Exception:
        return False


# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--input', required=True, help='folder with the Excel files')
    ap.add_argument('--base', default='data_base.json', help='last good data in the app schema (schema reference)')
    ap.add_argument('--current', default='data.json', help='current data.json; used for carry-forward if valid')
    ap.add_argument('--out', default='data.json')
    ap.add_argument('--site', default='ccr', choices=sorted(SITES), help='ccr = G0-D1810/D1801, d48 = G0-D4810/D4801')
    a = ap.parse_args()
    global SITE, CCR, SITE_RX
    SITE = a.site
    CCR, SITE_RX = SITES[SITE]
    log(f'site: {SITE} {CCR}')

    base = json.load(open(a.base, encoding='utf-8'))
    prev = base
    if a.current and os.path.exists(a.current):
        try:
            cur = json.load(open(a.current, encoding='utf-8'))
            if ok_schema(cur, base):
                prev = cur
                log(f'carry-forward source: {a.current} (data {cur["meta"].get("date")})')
            else:
                log(f'{a.current} is not in the app schema -> carry-forward from {a.base}')
        except Exception as e:
            log(f'{a.current} unreadable ({e}) -> carry-forward from {a.base}')
    prev = json.loads(json.dumps(prev))

    files = find_files(a.input)
    for k in PATTERNS:
        log(f'{k:7s}: {os.path.basename(files[k]) if k in files else "-- not found, carried forward"}')
    rows = {}
    for k, f in files.items():
        try:
            rows[k] = read_sheet(f, SHEETS[k])
        except Exception as e:
            log(f'  !! could not read {os.path.basename(f)}: {e} -> section carried forward')
    files = {k: f for k, f in files.items() if k in rows}   # sources that could not be read are not listed in meta

    out = {k: prev[k] for k in base}
    T = out['tabs']
    notes = []
    if rows.get('iwp'):
        out['items'] = build_items(rows['iwp'], prev, T)
    if rows.get('equip'):
        n, m = enrich_items_equipment(out['items'], rows['equip'], out['IK'])
        notes.append(f'equipment control: {m} CCR RFI records, {n} IWP items got RFI status in notes')
    prev_nc = len(prev['cables'])
    if rows.get('cables'):
        out['cables'], upd, add, rw = build_cables(rows['cables'], prev, out['items'], rows.get('rawp'))
        notes.append(f'cable tracker: {upd} cables updated, {add} new cables added (From/To kept for existing)')
        notes.append(f'RAWP: {rw} cables got laying/termination RFI info')
        out['panels'] = json.loads(json.dumps(prev['panels']))
        if not out['panels'] and SITE != 'ccr':
            out['panels'] = build_panels(out['cables'], out['CK'])
            prev_nc = len(out['cables'])
        notes.append(f'panels: {update_panels(out["panels"], out["cables"], prev_nc, out["CK"])} new cable links')
    if rows.get('master') or rows.get('rhi'):
        out['punch'], nm, nr = build_punch(rows.get('master'), rows.get('rhi'), prev)
        notes.append(f'punch: {nm} master (client) rows + {nr} RHI internal rows'
                     + (' (master kept from app)' if not rows.get('master') else '') + (' (RHI kept from app)' if not rows.get('rhi') else ''))
    if rows.get('docs'):
        out['reg'] = build_reg(rows['docs'], prev)
    if rows.get('ptq'):
        out['tq'] = dict(out['tq'], ptq=build_ptq(rows['ptq'], prev))
    if rows.get('ctq') and SITE == 'ccr':
        ctq, nreg = build_ctq(rows['ctq'], prev)
        out['tq'] = dict(out['tq'], ctq=ctq)
        notes.append(f'CTQ: {nreg} from register, {len(ctq) - nreg} carried forward')

    if SITE != 'ccr':      # D48: block diagrams rebuilt from the cable list every time
        out['sg'] = build_sg(out['cables'], out['CK'], out['panels'])
        notes.append(f'block diagrams: {len(out["sg"])} systems')
    # sg holds cable indices -> remap by tag if order changed
    old_t = [c[0] for c in prev['cables']]
    new_ix = {c[0]: i for i, c in enumerate(out['cables'])}
    for g in (out['sg'].values() if SITE == 'ccr' else []):
        g['e'] = [[e[0], e[1], [new_ix[old_t[i]] for i in e[2] if i < len(old_t) and old_t[i] in new_ix]] for e in g['e']]

    dates = [(file_key(f)[1], k) for k, f in files.items() if file_key(f)[1]]
    last = max(d for d, _ in dates) if dates else datetime.date.today()
    src = []
    for k, f in files.items():
        d = file_key(f)[1]
        src.append(f'{LABEL[k]} {d.strftime("%d.%m.%Y")}' if d else LABEL[k] + (' ' + re.search(r'_R(\d+)', f)[0][1:] if re.search(r'_R(\d+)', f) else ''))
    new_labels = [LABEL[k] for k in files]
    for part in str(prev['meta'].get('src', '')).split(' · '):
        if part and not any(part.startswith(l) for l in new_labels):
            src.append(part)
    try:
        pd = datetime.datetime.strptime(prev['meta'].get('date', ''), '%d.%m.%Y').date()
        last = max(last, pd)
    except ValueError:
        pass
    out['meta'] = {'date': last.strftime('%d.%m.%Y'), 'src': ' · '.join(src),
                   'ver': f'{last.isoformat()}-{datetime.datetime.now().strftime("%H%M")}'}
    if out['punch'] is not prev['punch'] and rows.get('master'):
        d = file_key(files['master'])[1]
        out['punch']['date'] = d.strftime('%d.%m.%Y') if d else out['meta']['date']

    out = scrub_all(out)
    notes.append(f'standard references removed from text: {STD_N[0]}')

    err = validate(out, base)
    log('\n'.join('note: ' + n for n in notes))
    log('\nsection        base   new   diff')
    def cnt(d):
        return {'items': len(d['items']), 'cables': len(d['cables']), 'panels': len(d['panels']),
                'punch.rows': len(d['punch']['rows']), 'drums': len(d['drums']), 'mat': len(d['mat']), 'mrf': len(d['mrf']),
                'reg': len(d['reg']), 'sf': len(d['sf']), 'sg': len(d['sg']), 'lr': len(d['lr']),
                'tq.ctq': len(d['tq']['ctq']), 'tq.ptq': len(d['tq']['ptq']), 'rv': len(d['rv'])}
    cb, cn = cnt(base), cnt(out)
    for k in cb:
        log(f'{k:12s} {cb[k]:6d} {cn[k]:6d} {cn[k] - cb[k]:+6d}')
    if err:
        log('\nSCHEMA VALIDATION FAILED - nothing written:\n  ' + '\n  '.join(err))
        sys.exit(1)
    tmp = a.out + '.tmp'
    with open(tmp, 'w', encoding='utf-8') as fh:
        json.dump(out, fh, ensure_ascii=False, separators=(',', ':'))
    os.replace(tmp, a.out)
    log(f'\nOK: wrote {a.out} ({os.path.getsize(a.out):,} bytes), meta {out["meta"]}')


if __name__ == '__main__':
    main()
